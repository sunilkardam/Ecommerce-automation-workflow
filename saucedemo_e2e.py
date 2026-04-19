"""
SauceDemo End-to-End Purchase Flow
Playwright (Python) — async

Dependencies:
    pip install playwright openpyxl
    playwright install chromium

Run:
    python saucedemo_e2e.py
"""

import asyncio
import os
import openpyxl
from playwright.async_api import async_playwright, Locator, Page, expect


# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
BASE_URL   = "https://www.saucedemo.com/"
EXCEL_FILE = os.path.join(os.path.dirname(__file__), "credentials.xlsx")
CUSTOMER   = {"first_name": "Sunil", "last_name": "Kardam", "zip": "410210"}

PRODUCTS = {
    "backpack": "Sauce Labs Backpack",
    "tshirt":   "Sauce Labs Bolt T-Shirt",
    "jacket":   "Sauce Labs Fleece Jacket",
}


# ---------------------------------------------------------------------------
# Read credentials from Excel
# ---------------------------------------------------------------------------
def load_credentials(path: str) -> dict:
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    headers = [str(cell.value).strip().lower() for cell in ws[1]]
    try:
        u_col = headers.index("username")
        p_col = headers.index("password")
    except ValueError as e:
        raise RuntimeError(
            f"credentials.xlsx must have 'Username' and 'Password' columns. {e}"
        )
    row = list(ws.iter_rows(min_row=2, values_only=True))[0]
    username = str(row[u_col]).strip()
    password = str(row[p_col]).strip()
    print(f"  [Excel] Loaded credentials for user: {username}")
    return {"username": username, "password": password}


# ---------------------------------------------------------------------------
# Cursor tracker — remembers where the mouse is between interactions
# ---------------------------------------------------------------------------
_cursor = {"x": 640.0, "y": 400.0}   # start at screen centre


async def smooth_move_to(page: Page, tx: float, ty: float,
                         steps: int = 40, delay_ms: int = 12) -> None:
    """
    Glide the cursor from its current tracked position to (tx, ty).
    Uses a smoothstep ease-in-out curve so the pointer accelerates
    away from the origin and decelerates as it arrives — just like a
    real hand moving a mouse.
    """
    sx, sy = _cursor["x"], _cursor["y"]
    for i in range(1, steps + 1):
        t = i / steps
        ease = t * t * (3.0 - 2.0 * t)          # smoothstep  0 → 1
        await page.mouse.move(
            sx + (tx - sx) * ease,
            sy + (ty - sy) * ease,
        )
        await page.wait_for_timeout(delay_ms)    # real delay between each step
    _cursor["x"] = tx
    _cursor["y"] = ty


async def move_to_element(page: Page, locator: Locator) -> tuple[float, float]:
    """Scroll element into view, compute its centre, glide the cursor there."""
    await locator.scroll_into_view_if_needed()
    box = await locator.bounding_box()
    cx  = box["x"] + box["width"]  / 2
    cy  = box["y"] + box["height"] / 2
    await smooth_move_to(page, cx, cy)
    return cx, cy


async def move_and_click(page: Page, locator: Locator, label: str = "") -> None:
    """
    Human-like interaction:
      1. Cursor travels from current position to the element centre
      2. Hovers for 1 second (visible pause on the button)
      3. Clicks
    """
    if label:
        print(f"  🖱  Moving to → {label}")
    cx, cy = await move_to_element(page, locator)
    await page.wait_for_timeout(1000)            # 1-second hover before click
    await page.mouse.click(cx, cy)


async def move_and_type(page: Page, locator: Locator, text: str, label: str = "") -> None:
    """
    Cursor travels to the input field, hovers 1 second, clicks to focus,
    then types character-by-character with a 120 ms keystroke delay.
    """
    if label:
        print(f"  🖱  Moving to → {label}")
    cx, cy = await move_to_element(page, locator)
    await page.wait_for_timeout(1000)
    await page.mouse.click(cx, cy)
    await locator.press_sequentially(text, delay=120)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def log(step: int, message: str) -> None:
    print(f"  [Step {step:02d}] {message}")


def section(title: str) -> None:
    print(f"\n{'─' * 55}")
    print(f"  {title}")
    print(f"{'─' * 55}")


def data_test_id(name: str) -> str:
    return name.lower().replace(" ", "-")


async def slow_scroll(page: Page, total_px: int = 500, steps: int = 10, delay_ms: int = 150) -> None:
    """Scroll down gradually in small increments."""
    step_px = total_px // steps
    for _ in range(steps):
        await page.mouse.wheel(0, step_px)
        await page.wait_for_timeout(delay_ms)


# ---------------------------------------------------------------------------
# Step 1 – Navigate and login
# ---------------------------------------------------------------------------
async def login(page: Page, creds: dict) -> None:
    section("STEP 1 · Navigate & Login")

    log(1, f"Navigating to {BASE_URL}")
    await page.goto(BASE_URL)
    await page.wait_for_load_state("domcontentloaded")
    await page.wait_for_timeout(1000)

    log(2, f"Typing username: {creds['username']}")
    await move_and_type(page, page.get_by_test_id("username"), creds["username"], "Username field")

    log(3, "Typing password")
    await move_and_type(page, page.get_by_test_id("password"), creds["password"], "Password field")

    print("  ⏳  Pausing 1 second before Login...")
    await page.wait_for_timeout(1000)

    log(4, "Clicking Login button")
    await move_and_click(page, page.get_by_test_id("login-button"), "Login button")

    log(5, "Validating inventory page")
    await page.wait_for_url(f"{BASE_URL}inventory.html")
    await expect(page.get_by_test_id("inventory-container")).to_be_visible()
    print("  ✔  Login successful — inventory page loaded")
    await page.wait_for_timeout(2000)


# ---------------------------------------------------------------------------
# Step 2 – Add Sauce Labs Backpack via product detail page
# ---------------------------------------------------------------------------
async def add_backpack(page: Page) -> None:
    section("STEP 2 · Add Sauce Labs Backpack")

    log(6, f"Clicking product title: {PRODUCTS['backpack']}")
    title_locator = (
        page.locator(".inventory_item")
        .filter(has_text=PRODUCTS["backpack"])
        .get_by_test_id("inventory-item-name")
    )
    await move_and_click(page, title_locator, f"'{PRODUCTS['backpack']}' title")

    await expect(page.get_by_test_id("inventory-item-name")).to_have_text(PRODUCTS["backpack"])
    print(f"  ✔  Product detail page loaded: {PRODUCTS['backpack']}")
    await page.wait_for_timeout(2000)

    log(7, "Clicking 'Add to cart'")
    await move_and_click(page, page.get_by_test_id("add-to-cart"), "Add to cart button")
    await expect(page.get_by_test_id("shopping-cart-badge")).to_have_text("1")
    print(f"  ✔  {PRODUCTS['backpack']} added — cart badge: 1")
    await page.wait_for_timeout(2000)

    log(8, "Clicking 'Back to products'")
    await move_and_click(page, page.get_by_test_id("back-to-products"), "Back to products button")
    await expect(page.get_by_test_id("inventory-container")).to_be_visible()
    print("  ✔  Back on inventory page")
    await page.wait_for_timeout(2000)


# ---------------------------------------------------------------------------
# Step 3 – Add Bolt T-Shirt and Fleece Jacket from inventory listing
# ---------------------------------------------------------------------------
async def add_inventory_items(page: Page) -> None:
    section("STEP 3 · Add T-Shirt & Fleece Jacket")

    for step_num, key in [(9, "tshirt"), (10, "jacket")]:
        name   = PRODUCTS[key]
        btn_id = f"add-to-cart-{data_test_id(name)}"
        log(step_num, f"Adding '{name}'")
        await move_and_click(page, page.get_by_test_id(btn_id), f"'Add to cart' — {name}")
        print(f"  ✔  {name} added to cart")
        await page.wait_for_timeout(2000)

    await expect(page.get_by_test_id("shopping-cart-badge")).to_have_text("3")
    print("  ✔  Cart badge updated to 3")


# ---------------------------------------------------------------------------
# Step 4 – Open cart, validate, scroll, remove Fleece Jacket, checkout
# ---------------------------------------------------------------------------
async def manage_cart(page: Page) -> None:
    section("STEP 4 · Cart Actions")

    log(11, "Clicking cart icon (top right)")
    await move_and_click(page, page.get_by_test_id("shopping-cart-link"), "Cart icon (top right)")
    await page.wait_for_url(f"{BASE_URL}cart.html")

    log(12, "Validating all 3 items are present in cart")
    cart_items = page.get_by_test_id("inventory-item-name")
    await expect(cart_items).to_have_count(3)
    item_names = await cart_items.all_inner_texts()
    print(f"  ✔  Cart contains: {item_names}")
    await page.wait_for_timeout(2000)

    assert PRODUCTS["backpack"] in item_names, "Backpack missing from cart"
    assert PRODUCTS["tshirt"]   in item_names, "T-Shirt missing from cart"
    assert PRODUCTS["jacket"]   in item_names, "Fleece Jacket missing from cart"

    print("  ↓  Scrolling down the cart page...")
    await slow_scroll(page)
    await page.wait_for_timeout(1000)

    log(13, f"Removing: {PRODUCTS['jacket']}")
    remove_locator = page.get_by_test_id(f"remove-{data_test_id(PRODUCTS['jacket'])}")
    await move_and_click(page, remove_locator, f"Remove '{PRODUCTS['jacket']}' button")

    log(14, "Validating Fleece Jacket is removed")
    await expect(cart_items).to_have_count(2)
    remaining = await cart_items.all_inner_texts()
    assert PRODUCTS["jacket"] not in remaining, "Fleece Jacket was NOT removed"
    print(f"  ✔  Fleece Jacket removed — cart now contains: {remaining}")
    await page.wait_for_timeout(2000)

    log(15, "Clicking 'Checkout'")
    await move_and_click(page, page.get_by_test_id("checkout"), "Checkout button")
    await page.wait_for_url(f"{BASE_URL}checkout-step-one.html")
    print("  ✔  Navigated to checkout information page")
    await page.wait_for_timeout(2000)


# ---------------------------------------------------------------------------
# Step 5 – Fill checkout information
# ---------------------------------------------------------------------------
async def checkout_info(page: Page) -> None:
    section("STEP 5 · Checkout Flow")

    log(16, "Filling customer information")
    await move_and_type(page, page.get_by_test_id("firstName"),  CUSTOMER["first_name"], "First Name field")
    await move_and_type(page, page.get_by_test_id("lastName"),   CUSTOMER["last_name"],  "Last Name field")
    await move_and_type(page, page.get_by_test_id("postalCode"), CUSTOMER["zip"],         "Zip Code field")
    print(
        f"  ✔  Entered: {CUSTOMER['first_name']} {CUSTOMER['last_name']}, "
        f"ZIP {CUSTOMER['zip']}"
    )
    await page.wait_for_timeout(2000)

    log(17, "Clicking 'Continue'")
    await move_and_click(page, page.get_by_test_id("continue"), "Continue button")
    await page.wait_for_url(f"{BASE_URL}checkout-step-two.html")
    print("  ✔  Checkout overview page loaded")


# ---------------------------------------------------------------------------
# Step 6 – Review overview, scroll, finish order
# ---------------------------------------------------------------------------
async def finish_order(page: Page) -> None:
    section("STEP 6 · Finish Order")

    log(18, "Validating checkout overview")
    await expect(page.get_by_test_id("checkout-summary-container")).to_be_visible()

    overview_items = page.get_by_test_id("inventory-item-name")
    await expect(overview_items).to_have_count(2)
    overview_names = await overview_items.all_inner_texts()
    print(f"  ✔  Overview items: {overview_names}")
    await page.wait_for_timeout(2000)

    print("  ↓  Scrolling down the overview page...")
    await slow_scroll(page, total_px=600, steps=12, delay_ms=150)
    await page.wait_for_timeout(1000)

    log(19, "Clicking 'Finish'")
    await move_and_click(page, page.get_by_test_id("finish"), "Finish button")
    await page.wait_for_url(f"{BASE_URL}checkout-complete.html")

    log(20, "Validating order confirmation")
    header = page.get_by_test_id("complete-header")
    await expect(header).to_be_visible()
    confirmation_text = await header.inner_text()
    assert "Thank you" in confirmation_text, (
        f"Unexpected confirmation text: '{confirmation_text}'"
    )
    print(f"  ✔  Order confirmed: \"{confirmation_text}\"")
    await page.wait_for_timeout(2000)


# ---------------------------------------------------------------------------
# Main orchestrator — single continuous flow
# ---------------------------------------------------------------------------
async def run_e2e_flow() -> None:
    print("\n" + "═" * 55)
    print("  SauceDemo · End-to-End Purchase Flow")
    print("═" * 55)

    creds = load_credentials(EXCEL_FILE)

    async with async_playwright() as pw:
        # SauceDemo uses data-test="…", not Playwright's default data-testid
        pw.selectors.set_test_id_attribute("data-test")

        browser = await pw.chromium.launch(headless=False)
        context = await browser.new_context(viewport={"width": 1280, "height": 800})
        page    = await context.new_page()

        try:
            await login(page, creds)
            await add_backpack(page)
            await add_inventory_items(page)
            await manage_cart(page)
            await checkout_info(page)
            await finish_order(page)

            print("\n" + "═" * 55)
            print("  ALL 20 STEPS PASSED ✔")
            print("═" * 55 + "\n")

        except Exception as exc:
            print(f"\n  ✘  TEST FAILED: {exc}")
            screenshot_path = os.path.join(os.path.dirname(__file__), "failure_screenshot.png")
            await page.screenshot(path=screenshot_path)
            print(f"  Screenshot saved → {screenshot_path}")
            raise

        finally:
            await context.close()
            await browser.close()


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------
if __name__ == "__main__":
    asyncio.run(run_e2e_flow())
